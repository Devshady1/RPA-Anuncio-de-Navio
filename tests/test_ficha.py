import os
import shutil
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import openpyxl

from anuncio_navio import ficha, historico

MODELO = Path(__file__).resolve().parents[1] / "src" / "anuncio_navio" / "recursos" / "ficha_modelo.xlsx"


class TestEstaSempre(unittest.TestCase):
    def test_ficha_recem_salva_nao_esta_livre(self):
        with TemporaryDirectory() as t:
            p = Path(t) / "f.xlsx"
            p.write_bytes(b"x")
            agora = p.stat().st_mtime + 1
            self.assertFalse(ficha.esta_livre(p, agora=agora))
            self.assertTrue(ficha.esta_livre(p, agora=agora + 60, lsof=lambda _p: ""))
            (Path(t) / ".~lock.f.xlsx#").write_text("lock")
            self.assertFalse(ficha.esta_livre(p, agora=agora + 60, lsof=lambda _p: ""))

    def test_lsof_aberto_bloqueia(self):
        with TemporaryDirectory() as t:
            p = Path(t) / "f.xlsx"
            p.write_bytes(b"x")
            self.assertFalse(ficha.esta_livre(p, agora=p.stat().st_mtime + 60, lsof=lambda _p: "123"))


class TestCicloDaFicha(unittest.TestCase):
    def test_escrever_status_preserva_o_resto(self):
        with TemporaryDirectory() as t:
            p = Path(t) / "f.xlsx"
            shutil.copyfile(MODELO, p)
            ficha.escrever_status(p, "NÃO FOI ENVIADO — teste")
            ws = openpyxl.load_workbook(p)["Anúncio"]
            self.assertEqual(ws["C3"].value, "NÃO FOI ENVIADO — teste")
            self.assertEqual(ws["A3"].value, "STATUS")

    def test_arquivar_move_com_duv(self):
        with TemporaryDirectory() as t:
            root = Path(t)
            (root / "ficha").mkdir()
            p = root / "ficha" / "ficha_anuncio_navio.xlsx"
            shutil.copyfile(MODELO, p)
            destino = ficha.arquivar(root, p, imo="9553256", viagem="000010", duv="000010")
            self.assertFalse(p.exists())
            self.assertEqual(destino.parent, root / "ficha" / "arquivo")
            self.assertIn("9553256_000010", destino.name)
            self.assertIn("DUV000010", destino.name)

    def test_gerar_nova_traz_historico(self):
        with TemporaryDirectory() as t:
            root = Path(t)
            (root / "ficha").mkdir()
            historico.registrar(root, navio="VOYAGER", imo="9553256", viagem="000009",
                                duv="000009", resultado="OK", detalhe="confirmado")
            nova = ficha.gerar_nova(root)
            ws = openpyxl.load_workbook(nova)["HISTÓRICO"]
            self.assertEqual([c.value for c in ws[1]], list(historico.CABECALHO))
            self.assertIn("000009", [c.value for c in ws[2]])

    def test_estado_evita_reprocessar(self):
        with TemporaryDirectory() as t:
            root = Path(t)
            (root / "ficha").mkdir()
            p = root / "ficha" / "ficha_anuncio_navio.xlsx"
            shutil.copyfile(MODELO, p)
            ficha.marcar_estado(root, p)
            self.assertTrue(ficha.ja_tentada(root, p))
            os.utime(p, (p.stat().st_atime, p.stat().st_mtime + 10))
            self.assertFalse(ficha.ja_tentada(root, p))


if __name__ == "__main__":
    unittest.main()
