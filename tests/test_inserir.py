import json
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import load_workbook  # noqa: E402

from inserir import ingerir_um  # noqa: E402

FICHA_MODELO = ROOT / "ficha_anuncio_navio.xlsx"

VALORES = {
    "nome do navio": "NAVIO TESTE",
    "imo": "9553256",
    "loa / comprimento (m)": "190.00",
    "último porto (nome ou código)": "Maceio",
    "eta · chegada": ("20/10/2026", "05:00"),
    "etb · atracação": ("21/10/2026", "06:00"),
    "etc · conclusão": ("22/10/2026", "12:00"),
    "ets · saída": ("22/10/2026", "18:00"),
    "próximo porto (nome ou código)": "Vitoria",
    "berço preferência": "101",
    "carga predominante (nome da mercadoria)": "niquel",
    "unidades (qtd_carga)": "1000",
    "tons previstas": "1000",
    "calado chegada (proa | popa)": ("9", "5"),
    "calado saída (proa | popa)": ("8", "4"),
    "agência (nome ou código)": "MARINAV",
    "armador (nome ou código)": "TAMAR SHIP MANAGEMENT",
    "operador (nome ou código)": "INTERMARIT",
}


def _preencher(destino: Path, limpar=()):
    wb = load_workbook(FICHA_MODELO)
    ws = wb.worksheets[0]
    rotulos = {str(ws.cell(row=r, column=1).value or "").strip().lower().rstrip("*").strip(): r
               for r in range(1, ws.max_row + 1)}
    for rot, v in VALORES.items():
        r = rotulos[rot.strip().lower()]
        if isinstance(v, tuple):
            ws.cell(row=r, column=3, value=v[0])
            ws.cell(row=r, column=4, value=v[1])
        else:
            ws.cell(row=r, column=3, value=v)
    for rot in limpar:
        ws.cell(row=rotulos[rot], column=3).value = None
    wb.save(destino)
    wb.close()


class TestInserir(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="inserir_test_"))
        self.entrada = self.tmp / "entrada"
        self.destino = self.tmp / "dados_rpa"
        self.seq = self.tmp / "seq"
        for d in (self.entrada, self.destino, self.seq / "data"):
            d.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def _ingerir(self, nome, limpar=()):
        arq = self.entrada / nome
        _preencher(arq, limpar=limpar)
        return ingerir_um(arq, destino=self.destino, root=self.seq)

    def test_ficha_valida_vira_json_com_o_texto_e_o_codigo_de_fallback(self):
        r = self._ingerir("valida.xlsx")
        self.assertTrue(r["ok"], r["erro"])
        payload = json.loads(r["json"].read_text(encoding="utf-8"))
        # O texto do agente vai intacto: quem resolve e o popup do OpenPort.
        self.assertEqual(payload["carga"], "niquel")
        self.assertEqual(payload["agencia"], "MARINAV")
        self.assertEqual(payload["operador"], "INTERMARIT")
        self.assertEqual(payload["armador"], "TAMAR SHIP MANAGEMENT")
        self.assertEqual(payload["porto_ant"], "Maceio")
        # O codigo do CSV local viaja ao lado, como fallback.
        self.assertEqual(payload["carga_codigo"], "26040000")
        self.assertEqual(payload["agencia_codigo"], "30010")
        self.assertEqual(payload["operador_codigo"], "30002")
        self.assertEqual(payload["armador_codigo"], "00000000000005")
        self.assertEqual(payload["porto_ant_codigo"], "BRMCZ")
        self.assertEqual(payload["tipo_viagem"], "1")   # classificacao padrao
        self.assertTrue(payload["viagem"])
        arquivada = list((self.entrada / "processados").glob("*.xlsx"))
        self.assertEqual(len(arquivada), 1)

    def test_ficha_invalida_nao_gera_json_nem_queima_viagem(self):
        r = self._ingerir("incompleta.xlsx", limpar=("imo", "carga predominante (nome da mercadoria)"))
        self.assertFalse(r["ok"])
        self.assertIsNone(r["json"])
        self.assertEqual(list(self.destino.glob("*.json")), [])
        self.assertFalse((self.seq / "data" / "sequencia.json").exists(),
                         "viagem reservada em ficha invalida!")
        errados = list((self.entrada / "erro").glob("*"))
        self.assertEqual(len(errados), 2, "esperava .xlsx + .log.txt")

    def test_ficha_travada_fica_na_entrada(self):
        arq = self.entrada / "travada.xlsx"
        _preencher(arq)
        handle = open(arq, "r+b")   # simula Excel com o arquivo aberto
        try:
            r = ingerir_um(arq, destino=self.destino, root=self.seq)
        finally:
            handle.close()
        self.assertTrue(r["travado"])
        self.assertIsNone(r["json"])
        self.assertEqual(list(self.destino.glob("*.json")), [])
        self.assertEqual(list(self.entrada.glob("*.xlsx")), [arq],
                         "ficha movida mesmo travada!")


if __name__ == "__main__":
    unittest.main()
