import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))

import openpyxl

from anuncio_navio import planilha

MODELO = Path(__file__).resolve().parents[1] / "src" / "anuncio_navio" / "recursos" / "ficha_modelo.xlsx"


class TestModeloFicha(unittest.TestCase):
    def test_abas_e_status_mesclado(self):
        wb = openpyxl.load_workbook(MODELO)
        self.assertEqual(wb.sheetnames, ["Anúncio", "COMO PREENCHER", "HISTÓRICO"])
        self.assertIn("C3:E3", [str(r) for r in wb["Anúncio"].merged_cells.ranges])

    def test_instrucao_de_salvar(self):
        ws = openpyxl.load_workbook(MODELO)["Anúncio"]
        self.assertIn("Salve", ws["A29"].value)

    def test_como_preencher_instrui_salvar_e_fechar(self):
        ws = openpyxl.load_workbook(MODELO)["COMO PREENCHER"]
        self.assertIn("salve (Cmd+S) e feche", ws["A9"].value)
        self.assertIn("automaticamente", ws["A9"].value)
        self.assertNotIn("devolva", ws["A9"].value.lower())
        self.assertIn("STATUS", ws["A10"].value)

    def test_leitor_reconhece_os_rotulos(self):
        campos = planilha.ler_planilha(MODELO)
        for chave in ("nome_navio", "imo", "eta_data", "carga", "agencia", "armador"):
            self.assertIn(chave, campos)
