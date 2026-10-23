"""Ciclo de vida da ficha unica: livre, resultado, arquivo, nova e estado.

A ficha e o unico arquivo que o agente ve. O servico so a processa quando
ela esta parada e fechada; o veredito volta na celula C3 da propria ficha;
a ficha preenchida vira evidencia em `ficha/arquivo/` e uma copia limpa do
modelo, com o HISTORICO mestre, ocupa o lugar dela.

    data/estado_ficha.json -> {"mtime": 123.45, "tamanho": 4567}

O estado evita reprocessar uma ficha parada: depois de gravar o veredito o
servico registra o novo `mtime` e a ficha so roda de novo se for salva.
"""
from __future__ import annotations

import json
import os
import shutil
import subprocess
import tempfile
import time
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

from . import caminhos, historico

MODELO = Path(__file__).resolve().parent / "recursos" / "ficha_modelo.xlsx"
ABA_ANUNCIO = "Anúncio"
ABA_HISTORICO = "HISTÓRICO"
ARQ_ESTADO = "data/estado_ficha.json"
CELULA_STATUS = "C3"
FORMATO_NOME = "%Y%m%d_%H%M%S"


def _lsof_padrao(path: Path) -> str:
    """Saida de `lsof -t <path>`; vazio quando o lsof nao existe."""
    try:
        saida = subprocess.run(
            ["lsof", "-t", str(path)],
            capture_output=True,
            text=True,
            timeout=10,
        )
    except (OSError, subprocess.SubprocessError):
        return ""
    return saida.stdout.strip()


def esta_livre(
    path: Path,
    *,
    agora: float | None = None,
    estavel_seg: float = 5.0,
    lsof=None,
) -> bool:
    """True se a ficha pode ser lida: mtime estavel, sem lock e sem lsof."""
    path = Path(path)
    try:
        mtime = os.stat(path).st_mtime
    except OSError:
        return False
    if mtime > (agora if agora is not None else time.time()) - estavel_seg:
        return False
    if any(True for _ in path.parent.glob(".~lock.*#")):
        return False
    if (lsof or _lsof_padrao)(path):
        return False
    return True


def _salvar_atomico(wb, destino: Path) -> None:
    fd, tmp = tempfile.mkstemp(dir=str(destino.parent), suffix=".xlsx")
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, destino)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise


def escrever_status(path: Path, mensagem: str) -> None:
    """Grava o veredito em C3 (vermelho) sem tocar no resto da ficha."""
    path = Path(path)
    wb = openpyxl.load_workbook(path)
    ws = wb[ABA_ANUNCIO]
    ws[CELULA_STATUS] = mensagem
    ws[CELULA_STATUS].font = Font(bold=True, color="9C0006")
    ws[CELULA_STATUS].fill = PatternFill("solid", fgColor="FFC7CE")
    ws[CELULA_STATUS].alignment = Alignment(wrap_text=True)
    _salvar_atomico(wb, path)


def arquivar(
    root: Path,
    origem: Path,
    *,
    imo: str,
    viagem: str,
    duv: str,
    quando=None,
) -> Path:
    """Move a ficha preenchida para `ficha/arquivo/` e devolve o destino."""
    momento = quando or datetime.now()
    arquivo_dir = caminhos.arquivo_dir(Path(root))
    arquivo_dir.mkdir(parents=True, exist_ok=True)
    nome = f"{imo}_{viagem}__{momento.strftime(FORMATO_NOME)}__DUV{duv}.xlsx"
    destino = arquivo_dir / nome
    shutil.move(str(origem), str(destino))
    return destino


def gerar_nova(root: Path) -> Path:
    """Copia o modelo, traz o HISTORICO mestre e o poe no lugar da ficha."""
    root = Path(root)
    destino = caminhos.ficha(root)
    destino.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=str(destino.parent), suffix=".xlsx")
    os.close(fd)
    try:
        shutil.copyfile(MODELO, tmp)
        wb = openpyxl.load_workbook(tmp)
        ws = wb[ABA_HISTORICO]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        linhas = list(historico.linhas(root))
        if linhas and tuple(linhas[0]) == historico.CABECALHO:
            linhas = linhas[1:]
        for numero, linha in enumerate(linhas, start=2):
            for coluna, valor in enumerate(linha, start=1):
                ws.cell(row=numero, column=coluna, value=valor)
        wb.save(tmp)
        os.replace(tmp, destino)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise
    return destino


def _caminho_estado(root: Path) -> Path:
    return Path(root) / ARQ_ESTADO


def marcar_estado(root: Path, path: Path) -> None:
    """Registra mtime e tamanho da ultima tentativa, de forma atomica."""
    st = os.stat(path)
    destino = _caminho_estado(Path(root))
    destino.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=str(destino.parent), suffix=".tmp")
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as f:
            json.dump({"mtime": st.st_mtime, "tamanho": st.st_size}, f, ensure_ascii=False, indent=2)
        os.replace(tmp, destino)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise


def ja_tentada(root: Path, path: Path) -> bool:
    """True se a ficha nao mudou desde a ultima tentativa registrada."""
    try:
        st = os.stat(path)
        dados = json.loads(_caminho_estado(Path(root)).read_text(encoding="utf-8"))
        return st.st_mtime == dados["mtime"] and st.st_size == dados["tamanho"]
    except (OSError, ValueError, TypeError, KeyError):
        return False
