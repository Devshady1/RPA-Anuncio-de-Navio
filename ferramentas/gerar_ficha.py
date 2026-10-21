# -*- coding: utf-8 -*-
"""Ficha de anuncio de navio v2 - linguagem visual do modelo oficial.

Nao toca no modelo_anuncio_navio.xlsx (oficial). Ordem dos campos na sequencia
em que o agente pensa na viagem; afretador aceita NOME (codigo sai sozinho).
Compativel com planilha.ler_planilha(): rotulos reconhecidos a partir da linha 2.
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "ficha_anuncio_navio.xlsx"

NAVY = "17365D"
HEAD = "244A73"
LABEL_BG = "EAF0F6"
INPUT_BG = "FFF5D8"
INPUT_OPT = "F5F7FA"
HINT_BG = "F7F9FB"
STAR = "A52A2A"
TXT = "23364A"
MUTED = "586B7C"

f_titulo = Font(name="Aptos Display", size=14, bold=True, color="FFFFFF")
f_sub = Font(name="Aptos", size=12, bold=True, color="DCE6F0")
f_status = Font(name="Aptos", size=10, bold=True, color=NAVY)
f_resp = Font(name="Aptos", size=10, bold=True, color="8A4B08")
f_cab = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
f_etapa = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
f_label = Font(name="Aptos", size=10, bold=True, color=NAVY)
f_star = Font(name="Aptos", size=11, bold=True, color=STAR)
f_hint = Font(name="Aptos", size=9, color=MUTED)
f_fim = Font(name="Aptos", size=10, bold=True, color="1F5C3A")

THIN = Side(style="thin", color="C9D2E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()
ws = wb.active
ws.title = "Anúncio"
ws.sheet_view.showGridLines = False
for col, w in (("A", 34), ("B", 5), ("C", 19), ("D", 19), ("E", 67)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A5"

# (tipo, rotulo, obrigatorio, duplo, dica)
LINHAS = [
    ("ETAPA", "ETAPA 1 — O NAVIO", 0, 0, ""),
    ("CAMPO", "Nome do Navio", 1, 0, "Nome oficial da embarcação (como está no documento do navio)"),
    ("CAMPO", "IMO", 1, 0, "7 dígitos numéricos, sem espaços"),
    ("CAMPO", "LOA / Comprimento (m)", 1, 0, "Comprimento total do navio, em metros"),
    ("CAMPO", "TPB / DWT", 0, 0, "Opcional — porte bruto, em toneladas"),
    ("ETAPA", "ETAPA 2 — A VIAGEM (de onde vem, quando chega e para onde vai)", 0, 0, ""),
    ("CAMPO", "Último Porto (nome ou código)", 1, 0, "Porto de origem — pode escrever o NOME da cidade"),
    ("DUPLO", "ETA · chegada", 1, 1, "C = data (dd/mm/aaaa); D = hora (hh:mm)"),
    ("DUPLO", "ETB · atracação", 1, 1, "C = data; D = hora. Depois da ETA"),
    ("DUPLO", "ETC · conclusão", 1, 1, "C = data; D = hora. Depois da ETB"),
    ("DUPLO", "ETS · saída", 1, 1, "C = data; D = hora. A última da viagem"),
    ("CAMPO", "Próximo Porto (nome ou código)", 1, 0, "Porto de destino — pode escrever o NOME da cidade"),
    ("CAMPO", "Berço preferência", 0, 0, "Opcional — só se houver preferência de berço"),
    ("ETAPA", "ETAPA 3 — A CARGA E OS CALADOS", 0, 0, ""),
    ("CAMPO", "Carga Predominante (nome da mercadoria)", 1, 0, "Nome da mercadoria — o código NCM é resolvido automaticamente"),
    ("CAMPO", "Unidades (QTD_CARGA)", 1, 0, "Quantidade de unidades (0 se não houver)"),
    ("CAMPO", "Tons previstas", 0, 0, "Opcional — toneladas previstas"),
    ("DUPLO", "Calado chegada (PROA | POPA)", 1, 1, "C = PROA; D = POPA. Como o navio CHEGA, em metros (até 16,0)"),
    ("DUPLO", "Calado saída (PROA | POPA)", 1, 1, "C = PROA; D = POPA. Como o navio SAI, em metros (até 16,0)"),
    ("ETAPA", "ETAPA 4 — OS RESPONSÁVEIS (pode escrever o NOME — o código sai sozinho)", 0, 0, ""),
    ("CAMPO", "Agência (nome ou código)", 1, 0, "Agência marítima — nome ou código"),
    ("CAMPO", "Armador (nome ou código)", 1, 0, "Armador — nome ou código"),
    ("CAMPO", "Operador (nome ou código)", 1, 0, "Operador — nome ou código"),
    ("CAMPO", "Afretador (nome ou código)", 0, 0, "Opcional — só se houver afretador; nome ou código"),
    ("FIM", "✔  Preencheu? Devolva este arquivo — não apague nem altere os textos da coluna A", 0, 0, ""),
]

rows_inputs = []   # celulas de preenchimento p/ formula de STATUS
ws.append([None, None, "AUTORIDADE PORTUÁRIA CODEBA", None, None])
ws.append([None, None, "FICHA DE ANÚNCIO DE NAVIO", None, None])
ws.append(["STATUS", None, "", None, None])

r = 4
ws.cell(row=r, column=1, value="CAMPO")
ws.cell(row=r, column=2, value="★")
ws.cell(row=r, column=3, value="VALOR")
ws.cell(row=r, column=4, value="VALOR 2 / HORA")
ws.cell(row=r, column=5, value="COMO PREENCHER")
for ci in range(1, 6):
    c = ws.cell(row=r, column=ci)
    c.font, c.fill, c.alignment = f_cab, PatternFill("solid", fgColor=HEAD), CENTER

for tipo, rotulo, obrig, duplo, dica in LINHAS:
    r += 1
    ws.cell(row=r, column=1, value=rotulo)
    ws.cell(row=r, column=2, value="★" if obrig else None)
    ws.cell(row=r, column=5, value=dica or None)
    for ci in range(1, 6):
        c = ws.cell(row=r, column=ci)
        c.border = BORDER
        c.alignment = LEFT if ci in (1, 5) else CENTER
    if tipo == "ETAPA":
        for ci in range(1, 6):
            c = ws.cell(row=r, column=ci)
            c.fill = PatternFill("solid", fgColor=NAVY)
            c.font = f_etapa
            c.alignment = LEFT
        continue
    if tipo == "FIM":
        for ci in range(1, 6):
            c = ws.cell(row=r, column=ci)
            c.fill = PatternFill("solid", fgColor="E8F2EB")
            c.font = f_fim
        continue
    ws.cell(row=r, column=1).font = f_label
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=LABEL_BG)
    if obrig:
        ws.cell(row=r, column=2).font = f_star
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=LABEL_BG)
        ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=INPUT_BG)
        rows_inputs.append(f"C{r}")
        if duplo:
            ws.cell(row=r, column=4).fill = PatternFill("solid", fgColor=INPUT_BG)
            rows_inputs.append(f"D{r}")
    else:
        ws.cell(row=r, column=2).font = Font(name="Aptos", size=11, color=MUTED)
        ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=LABEL_BG)
        ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=INPUT_OPT)
    ws.cell(row=r, column=5).font = f_hint
    ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=HINT_BG)

# STATUS (igual ao do modelo oficial: formula viva de preenchimento)
partes = "+".join(f"COUNTBLANK({c})" for c in rows_inputs)
ws["C3"] = f'=IF(({partes})=0,"OK — PRONTO PARA ENVIAR","FALTAM "&({partes})&" CAMPO(S)")'
ws["A3"].font, ws["A3"].fill = f_status, PatternFill("solid", fgColor=LABEL_BG)
ws["C3"].font, ws["C3"].fill = f_resp, PatternFill("solid", fgColor="FFF1CC")
ws["C3"].alignment, ws["A3"].alignment = CENTER, CENTER

for rr, (font, fill, align) in {
    1: (f_titulo, NAVY, CENTER),
    2: (f_sub, NAVY, CENTER),
}.items():
    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=5)
    c = ws.cell(row=rr, column=1)
    c.font, c.alignment = font, align
    for ci in range(1, 6):
        ws.cell(row=rr, column=ci).fill = PatternFill("solid", fgColor=fill)

ws2 = wb.create_sheet("COMO PREENCHER")
ws2.sheet_view.showGridLines = False
ws2.column_dimensions["A"].width = 96
instrucoes = [
    ("COMO PREENCHER ESTA FICHA", True),
    ("Preencha SOMENTE as células claras da coluna C (e da coluna D nos campos de DATA/HORA e CALADO).", False),
    ("Não apague nem altere os textos da coluna A — é por eles que o robô encontra cada dado.", False),
    ("Os campos com ★ são obrigatórios. A linha STATUS avisa quando faltar algum.", False),
    ("Agência, armador, operador, afretador e portos: pode escrever o NOME. O código é resolvido automaticamente.", False),
    ("Carga: escreva o NOME da mercadoria. O código NCM é resolvido automaticamente; "
     "se o nome for ambíguo, o RPA mostra as opções para escolha.", False),
    ("Calados em metros (até 16,0). Chegada = como o navio chega; saída = como o navio sai.", False),
    ("Datas: dia/mês/ano e hora. A ordem é ETA, depois ETB, ETC e por último ETS.", False),
    ("Pronto: devolva ESTE arquivo. O número da viagem e a classificação da operação são "
     "definidos no processamento.", False),
]
for i, (txt, bold) in enumerate(instrucoes, start=1):
    c = ws2.cell(row=i, column=1, value=txt)
    c.font = Font(name="Aptos", size=12 if i == 1 else 10, bold=bold,
                  color=NAVY if bold else TXT)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    if bold:
        for ci in range(1, 2):
            ws2.cell(row=i, column=ci).fill = PatternFill("solid", fgColor=LABEL_BG)

wb.save(OUT)
print("Ficha v2 OK:", OUT, OUT.stat().st_size, "bytes",
      "| linhas:", ws.max_row, "| obrigatorios:", len(rows_inputs))
