# -*- coding: utf-8 -*-
"""Gera relatorio institucional sobre a automacao do anuncio de navios."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

ROOT = Path(__file__).resolve().parent.parent
OUT_DIR = ROOT / "relatorios"
PDF_OUT = OUT_DIR / "NOTA_GERENCIAL_AUTOMACAO_ANUNCIO_NAVIOS_CODEBA_20261026.pdf"
DOCX_OUT = OUT_DIR / "NOTA_GERENCIAL_AUTOMACAO_ANUNCIO_NAVIOS_CODEBA_20261026.docx"
LOGO = ROOT / "logo_codeba_transp.png"
AS_OF = "26 de outubro de 2026"

NAVY_HEX = "12304A"
BLUE_HEX = "1E5875"
GOLD_HEX = "C7A451"
INK_HEX = "243746"
MUTED_HEX = "617382"
PALE_HEX = "F2F6F8"
LINE_HEX = "D9E2E8"
GREEN_HEX = "19734A"
GREEN_PALE_HEX = "EAF5EF"
AMBER_HEX = "8A5A08"
AMBER_PALE_HEX = "FFF5E2"

NAVY = colors.HexColor(f"#{NAVY_HEX}")
BLUE = colors.HexColor(f"#{BLUE_HEX}")
GOLD = colors.HexColor(f"#{GOLD_HEX}")
INK = colors.HexColor(f"#{INK_HEX}")
MUTED = colors.HexColor(f"#{MUTED_HEX}")
PALE = colors.HexColor(f"#{PALE_HEX}")
LINE = colors.HexColor(f"#{LINE_HEX}")
GREEN = colors.HexColor(f"#{GREEN_HEX}")
GREEN_PALE = colors.HexColor(f"#{GREEN_PALE_HEX}")
AMBER = colors.HexColor(f"#{AMBER_HEX}")
AMBER_PALE = colors.HexColor(f"#{AMBER_PALE_HEX}")
WHITE = colors.white

REPORTLAB_FONTS = Path(__import__("reportlab").__file__).resolve().parent / "fonts"
pdfmetrics.registerFont(TTFont("MgmtSans", str(REPORTLAB_FONTS / "Vera.ttf")))
pdfmetrics.registerFont(TTFont("MgmtSans-Bold", str(REPORTLAB_FONTS / "VeraBd.ttf")))

PAGES = [
    {
        "eyebrow": "01 · Síntese para decisão",
        "title": "Uma rotina portuária mais simples, controlada e rastreável",
        "blocks": [
            ("lead", "Esta iniciativa organiza e automatiza o preenchimento do anúncio de navios no OpenPort da CODEBA. A proposta é apoiar o agente marítimo na preparação das informações e retirar dele a necessidade de navegar pelo sistema e conhecer códigos internos de cadastro."),
            ("callout", ("Em uma frase", "O agente informa os dados da escala em uma ficha orientada; o serviço valida o preenchimento, registra o anúncio no OpenPort e confere o resultado na consulta operacional antes de devolver o status e o número da requisição (Programação) e arquivar a ficha.")),
            ("heading", "Por que isso interessa à gestão portuária"),
            ("body", "O anúncio é uma etapa administrativa relevante para que as informações da escala estejam disponíveis no sistema utilizado pela operação. A automação busca tornar essa etapa mais uniforme e verificável, com menos dependência de navegação manual e de conhecimento de códigos de cadastro."),
            ("cards", [
                ("Público diretamente apoiado", "Agente marítimo"),
                ("Área interlocutora", "Operação portuária · GPI"),
                ("Sistema atendido", "OpenPort / CODEBA"),
            ]),
            ("heading", "Mensagem executiva"),
            ("body", "Há uma prova anterior de anúncio realizado em produção e, separadamente, uma validação recente do fluxo de ficha única no QAS. Isso demonstra viabilidade técnica em etapas distintas; não significa que o serviço já esteja instalado ou autorizado para uso rotineiro no equipamento do agente."),
            ("callout_warning", ("Ponto de transparência", "O pacote de instalação para o notebook do agente ainda depende de identificação do modelo e sistema operacional, escolha do aplicativo de planilha e validação no equipamento real. A operação assistida em produção não deve ser confundida com a implantação operacional da ficha única.")),
        ],
    },
    {
        "eyebrow": "02 · Necessidade e público atendido",
        "title": "O processo começa com a informação da escala",
        "blocks": [
            ("heading", "A necessidade observada"),
            ("body", "O agente marítimo é responsável por informar dados essenciais da embarcação e da viagem. Para registrar o anúncio diretamente, também precisa localizar telas, preencher campos e selecionar cadastros do OpenPort. Alguns desses cadastros são identificados por códigos que não são necessariamente conhecidos pelo agente."),
            ("body", "A solução em desenvolvimento separa essas responsabilidades: o agente informa o que conhece sobre a escala; a ferramenta organiza os dados, valida campos básicos e auxilia a seleção dos cadastros existentes. A decisão operacional sobre os dados continua sendo humana."),
            ("heading", "Quem é apoiado e de que forma"),
            ("table", [
                ("Agente marítimo", "Recebe uma ficha guiada; informa os dados da escala; salva e fecha o arquivo. Não precisa executar a navegação do anúncio no OpenPort."),
                ("GPI / área operacional", "Pode acompanhar um fluxo mais padronizado, consultar o resultado e identificar fichas com dados pendentes. A definição formal de responsabilidades e do procedimento ainda deve ser acordada."),
                ("Gestão portuária", "Pode ter maior visibilidade sobre a existência, resultado e pendências do processo, com registros de histórico. Não se afirma ainda ganho financeiro, redução de prazo ou impacto de produtividade medido."),
                ("TI / suporte", "Mantém o ambiente de execução, credenciais protegidas, atualizações e suporte técnico, após a definição do modelo de implantação."),
            ]),
            ("heading", "O que não muda"),
            ("bullets", [
                "O agente ou a área responsável continua sendo a fonte dos dados comerciais e operacionais da escala.",
                "A automação não substitui a validação da informação por quem conhece a operação.",
                "O OpenPort continua sendo o sistema de registro; a ferramenta opera a interface e confere o resultado.",
                "A ficha não deve ser processada enquanto estiver aberta ou sendo editada.",
            ]),
            ("callout", ("Benefício esperado, ainda a medir", "Mais consistência no preenchimento e melhor rastreabilidade do ciclo. O tamanho do benefício deverá ser avaliado com indicadores do piloto, não presumido.")),
        ],
    },
    {
        "eyebrow": "03 · Como o fluxo proposto funciona",
        "title": "Da ficha ao registro conferido",
        "blocks": [
            ("steps", [
                ("1", "Preencher", "O agente marítimo informa os dados da embarcação, rota, horários previstos, carga e demais campos da ficha."),
                ("2", "Validar", "O serviço aguarda o arquivo ser salvo e fechado, verifica campos obrigatórios e aponta o que precisa ser corrigido."),
                ("3", "Registrar", "Com dados válidos, o serviço acessa o OpenPort e preenche o anúncio, apoiando a seleção dos cadastros disponíveis."),
                ("4", "Conferir", "Depois da gravação, consulta a tela operacional de programação para localizar o registro e conferir os dados de identificação e status."),
                ("5", "Retornar", "Atualiza o resultado na ficha, registra histórico e evidências, arquiva a ficha processada e prepara uma nova ficha."),
            ]),
            ("heading", "Como são tratados códigos e cadastros"),
            ("body", "Quando a ficha contém nomes conhecidos pela operação — como porto, agência ou operador — o resolvedor consulta as opções do OpenPort para relacionar a descrição ao cadastro correto. Tabelas locais apoiam o processo quando aplicável. Se uma correspondência não for segura ou houver ambiguidade, o comportamento esperado é interromper e sinalizar, em vez de escolher um cadastro por tentativa."),
            ("heading", "Resultado não é presumido"),
            ("body", "A gravação no formulário, por si só, não é tratada como sucesso final. O fluxo de serviço considera concluído quando a verificação posterior encontra o anúncio esperado. Se houver erro de preenchimento, acesso ou conferência, a ficha recebe um retorno para correção ou análise."),
            ("callout_warning", ("Regra de operação", "Apenas um anúncio pode ser processado por vez. Uma segunda tentativa simultânea é recusada; a área deve aguardar a conclusão ou o retorno do primeiro ciclo.")),
            ("heading", "Automação assistida, não decisão autônoma"),
            ("body", "O software executa uma sequência de tarefas repetitivas na interface. Não escolhe a escala, não negocia dados e não substitui a aprovação operacional. Dados faltantes, inconsistências ou mudanças de condição devem voltar à pessoa responsável."),
        ],
    },
    {
        "eyebrow": "04 · O que foi realizado e comprovado",
        "title": "Duas validações distintas, com registros verificáveis",
        "blocks": [
            ("table", [
                ("13/10/2026 · Produção", "TI SUNRISE, IMO 9701009, viagem/DUV 000006. Anúncio confirmado e localizado na consulta 2001 como programação 1772, status “Anunciado”. A execução utilizou o fluxo de fila JSON já existente."),
                ("26/10/2026 · QAS", "VOYAGER, IMO 9553256, viagem/DUV 009998, ETA 28/12/2026. Uma execução ponta a ponta do serviço de ficha única foi confirmada na consulta 2001 como programação 1782, status “Anunciado”."),
                ("26/10/2026 · Testes locais", "A suíte automatizada passou com 143 testes; 4 verificações opt-in que poderiam usar QAS ficaram ignoradas nessa execução. Os testes locais cobriram configuração, bloqueios de ambiente, trava e fluxo do serviço."),
            ]),
            ("heading", "O que esses resultados permitem afirmar"),
            ("bullets", [
                "A automação já foi capaz de concluir um anúncio real em produção, no fluxo anterior baseado em dados JSON.",
                "O fluxo mais recente de ficha única foi executado de ponta a ponta uma vez no QAS, em raiz temporária isolada.",
                "A verificação independente na consulta 2001 foi usada nos dois registros descritos.",
                "Foi implementada trava para impedir anúncios simultâneos e controles para restringir o ambiente de execução.",
            ]),
            ("heading", "O que ainda não permitem afirmar"),
            ("bullets", [
                "A ficha única ainda não foi instalada nem validada no notebook do agente.",
                "O fluxo de ficha única ainda não teve uma operação rotineira em produção.",
                "Não há medição aprovada de redução de tempo, custo, retrabalho ou impacto sobre a capacidade do porto.",
                "O caso de produção de 13/10 e a validação de ficha única no QAS são evidências diferentes; uma não substitui a outra.",
            ]),
            ("callout", ("Rastreabilidade", "O projeto mantém histórico, logs e evidências de execução. O teste QAS de 26/10 foi registrado; seu identificador de viagem/ETA não deve ser reutilizado.")),
        ],
    },
    {
        "eyebrow": "05 · Controles, responsabilidades e limites",
        "title": "Confiabilidade depende de controles técnicos e operacionais",
        "blocks": [
            ("table", [
                ("Dados de entrada", "Campos obrigatórios e formatos são validados antes de iniciar o anúncio. Informações operacionais não devem ser inferidas sem confirmação."),
                ("Cadastros", "O sistema utiliza os cadastros do OpenPort e trata ambiguidades como pendência; a área responsável valida os dados fornecidos."),
                ("Confirmação", "A conclusão do serviço exige localização do registro esperado na consulta operacional, não apenas o clique em gravar."),
                ("Concorrência", "Uma trava operacional permite somente um anúncio ativo por vez, inclusive entre o serviço e execuções iniciadas manualmente."),
                ("Ambiente", "Durante o desenvolvimento, QAS é o padrão; produção só é acessada quando solicitada. URLs desconhecidas são bloqueadas."),
                ("Credenciais", "As credenciais não devem ser incluídas no repositório. A preparação para macOS prevê armazenamento de senha no Keychain."),
                ("Falhas", "O retorno deve indicar que o anúncio não foi enviado, o motivo conhecido e a ação de correção ou análise."),
            ]),
            ("heading", "Responsabilidades a pactuar antes da rotina"),
            ("bullets", [
                "Agente marítimo: preencher a ficha com dados confirmados e mantê-la salva e fechada para processamento.",
                "GPI / operação: definir quem revisa pendências, quem acompanha o retorno e como se procede em alterações da escala.",
                "TI / responsável pelo serviço: instalar, proteger credenciais, acompanhar disponibilidade, logs e atualizações.",
                "Gestão: aprovar a entrada em operação, as condições de uso em produção e os indicadores de acompanhamento.",
            ]),
            ("body_small", "Os papéis acima são uma proposta de governança para discussão; não representam designações formais já aprovadas."),
            ("callout_warning", ("Limite importante", "A automação depende de disponibilidade do OpenPort, rede, credenciais válidas e compatibilidade da interface. Erros no portal ou mudanças de tela podem exigir intervenção e manutenção.")),
        ],
    },
    {
        "eyebrow": "06 · Próximas etapas e acompanhamento",
        "title": "Do piloto validado à decisão de implantação",
        "blocks": [
            ("steps", [
                ("1", "Identificar o equipamento", "Confirmar modelo completo e sistema operacional do notebook “Neo”, arquitetura e aplicativo de planilha."),
                ("2", "Preparar a entrega", "Construir instalador, runtime do navegador, configuração segura de credenciais e inicialização controlada."),
                ("3", "Validar no equipamento", "Testar abertura e fechamento da ficha, notificações, logs, recuperação de falhas e regra de um anúncio por vez."),
                ("4", "Autorizar a operação", "Definir responsáveis, suporte, contingência e aprovação explícita para produção no equipamento operacional."),
                ("5", "Medir o piloto", "Acompanhar resultados por período acordado antes de concluir sobre produtividade ou expansão."),
            ]),
            ("heading", "Indicadores recomendados para o piloto"),
            ("table", [
                ("Volume", "Fichas recebidas e anúncios concluídos por período."),
                ("Qualidade", "Percentual concluído sem correção; tipos de pendência e retrabalho."),
                ("Tempo", "Tempo entre ficha pronta e retorno do resultado, distinguindo espera do portal e correção humana."),
                ("Continuidade", "Indisponibilidades, falhas de acesso e necessidade de intervenção."),
                ("Controle", "Ocorrências de ficha aberta, duplicidade detectada e tentativa simultânea recusada."),
            ]),
            ("heading", "Decisões solicitadas à gestão"),
            ("bullets", [
                "Designar interlocutor operacional do GPI e responsável pelo acompanhamento do piloto.",
                "Confirmar limites do piloto: usuários, período, procedimento de contingência e canal de suporte.",
                "Autorizar a configuração para produção somente após validação no equipamento e aceite operacional.",
                "Definir quais indicadores e evidências serão apresentados ao final do período de avaliação.",
            ]),
            ("callout", ("Conclusão", "A iniciativa tem evidência de viabilidade e uma trilha de implantação ainda em aberto. O próximo passo prudente é concluir a validação no equipamento, pactuar responsabilidades e acompanhar um piloto controlado antes de avaliar ganhos ou ampliar o uso.")),
            ("body_small", "Base de elaboração: documentação operacional e técnica do projeto, registros de validação e evidências existentes até 26/10/2026. Documento informativo para discussão gerencial; não substitui norma, procedimento portuário ou autorização formal."),
        ],
    },
]


def pdf_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "MgmtTitle", parent=base["Title"], fontName="MgmtSans-Bold",
            fontSize=25, leading=30, textColor=NAVY, alignment=TA_LEFT,
            spaceAfter=3 * mm,
        ),
        "subtitle": ParagraphStyle(
            "MgmtSubtitle", parent=base["Normal"], fontName="MgmtSans",
            fontSize=10, leading=14, textColor=MUTED, spaceAfter=5 * mm,
        ),
        "eyebrow": ParagraphStyle(
            "MgmtEyebrow", parent=base["Normal"], fontName="MgmtSans-Bold",
            fontSize=8, leading=10, textColor=BLUE, spaceAfter=2 * mm,
        ),
        "heading": ParagraphStyle(
            "MgmtHeading", parent=base["Heading2"], fontName="MgmtSans-Bold",
            fontSize=12, leading=15, textColor=NAVY, spaceBefore=3 * mm,
            spaceAfter=2 * mm, keepWithNext=True,
        ),
        "body": ParagraphStyle(
            "MgmtBody", parent=base["BodyText"], fontName="MgmtSans",
            fontSize=9, leading=13, textColor=INK, spaceAfter=2.3 * mm,
        ),
        "lead": ParagraphStyle(
            "MgmtLead", parent=base["BodyText"], fontName="MgmtSans",
            fontSize=11, leading=16, textColor=INK, spaceAfter=4 * mm,
        ),
        "small": ParagraphStyle(
            "MgmtSmall", parent=base["BodyText"], fontName="MgmtSans",
            fontSize=7.7, leading=10.5, textColor=MUTED, spaceAfter=2 * mm,
        ),
        "label": ParagraphStyle(
            "MgmtLabel", parent=base["BodyText"], fontName="MgmtSans-Bold",
            fontSize=7.5, leading=10, textColor=BLUE,
        ),
        "cell": ParagraphStyle(
            "MgmtCell", parent=base["BodyText"], fontName="MgmtSans",
            fontSize=8, leading=11, textColor=INK,
        ),
        "step": ParagraphStyle(
            "MgmtStep", parent=base["BodyText"], fontName="MgmtSans",
            fontSize=8, leading=11, textColor=INK,
        ),
    }


def pdf_table(rows, widths, *, shaded_first=False):
    prepared = [
        [Paragraph(str(left), pdf_styles_cache["label"]), Paragraph(str(right), pdf_styles_cache["cell"])]
        for left, right in rows
    ]
    table = Table(prepared, colWidths=widths, hAlign="LEFT")
    commands = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.45, LINE),
        ("LEFTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 2.5 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 2 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2 * mm),
    ]
    if shaded_first:
        commands.append(("BACKGROUND", (0, 0), (0, -1), PALE))
    table.setStyle(TableStyle(commands))
    return table


def pdf_callout(title_text, body_text, warning=False):
    fill = AMBER_PALE if warning else GREEN_PALE
    accent = AMBER if warning else GREEN
    style = ParagraphStyle(
        f"Callout{len(title_text)}{warning}", parent=pdf_styles_cache["body"],
        spaceAfter=0,
    )
    box = Table([[Paragraph(f"<b>{title_text}</b><br/>{body_text}", style)]], colWidths=[180 * mm])
    box.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), fill),
        ("LINEBEFORE", (0, 0), (0, 0), 3, accent),
        ("LEFTPADDING", (0, 0), (-1, -1), 4 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 3 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3 * mm),
    ]))
    return box


def pdf_bullets(items):
    result = []
    for item in items:
        result.append(Paragraph(f"•  {item}", pdf_styles_cache["body"]))
    return result


def pdf_steps(items):
    rows = []
    for number, heading, detail in items:
        text = f"<font color='#{BLUE_HEX}'><b>{number} · {heading}</b></font><br/>{detail}"
        rows.append([Paragraph(text, pdf_styles_cache["step"])])
    table = Table(rows, colWidths=[180 * mm], hAlign="LEFT")
    table.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LINEBEFORE", (0, 0), (0, -1), 2, GOLD),
        ("LINEBELOW", (0, 0), (-1, -1), 0.45, LINE),
        ("LEFTPADDING", (0, 0), (-1, -1), 4 * mm),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm),
        ("TOPPADDING", (0, 0), (-1, -1), 2.3 * mm),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.3 * mm),
    ]))
    return table


def pdf_page_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(GOLD)
    canvas.setLineWidth(0.7)
    canvas.line(15 * mm, 15 * mm, A4[0] - 15 * mm, 15 * mm)
    canvas.setFont("MgmtSans", 7)
    canvas.setFillColor(MUTED)
    canvas.drawString(15 * mm, 10 * mm, f"CODEBA · GPI / ÁREA OPERACIONAL · {AS_OF.upper()}")
    canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"{doc.page}  /  {len(PAGES)}")
    canvas.restoreState()


def build_pdf():
    global pdf_styles_cache
    pdf_styles_cache = pdf_styles()
    story = []
    for page_number, page in enumerate(PAGES, start=1):
        if page_number == 1:
            logo = Image(str(LOGO), width=46 * mm, height=46 * mm * 688 / 1555)
            header = Table(
                [[logo, Paragraph(
                    "<b>GESTÃO PORTUÁRIA · GPI</b>",
                    ParagraphStyle(
                        "Brand", parent=pdf_styles_cache["body"], fontName="MgmtSans-Bold",
                        textColor=WHITE, alignment=TA_LEFT, fontSize=10, leading=14,
                    ),
                )]],
                colWidths=[57 * mm, 123 * mm],
                rowHeights=[27 * mm],
            )
            header.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, -1), NAVY),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LINEBELOW", (0, 0), (-1, -1), 1.5, GOLD),
                ("LEFTPADDING", (0, 0), (-1, -1), 3 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 2 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2 * mm),
            ]))
            story.extend([header, Spacer(1, 9 * mm)])
            story.append(Paragraph("AUTOMAÇÃO DO ANÚNCIO<br/>DE NAVIOS NO OPENPORT", pdf_styles_cache["title"]))
            story.append(Paragraph(
                "Visão gerencial e operacional · Documento informativo para discussão",
                pdf_styles_cache["subtitle"],
            ))
            story.append(Spacer(1, 2 * mm))
            story.append(Paragraph(page["eyebrow"], pdf_styles_cache["eyebrow"]))
        else:
            story.append(Paragraph(page["eyebrow"], pdf_styles_cache["eyebrow"]))
            story.append(Paragraph(page["title"], pdf_styles_cache["title"]))

        for kind, value in page["blocks"]:
            if kind == "heading":
                story.append(Paragraph(value, pdf_styles_cache["heading"]))
            elif kind in ("body", "lead", "body_small"):
                style = {
                    "body": pdf_styles_cache["body"],
                    "lead": pdf_styles_cache["lead"],
                    "body_small": pdf_styles_cache["small"],
                }[kind]
                story.append(Paragraph(value, style))
            elif kind == "callout":
                story.append(pdf_callout(value[0], value[1]))
                story.append(Spacer(1, 2 * mm))
            elif kind == "callout_warning":
                story.append(pdf_callout(value[0], value[1], warning=True))
                story.append(Spacer(1, 2 * mm))
            elif kind == "table":
                story.append(pdf_table(value, [47 * mm, 133 * mm], shaded_first=True))
                story.append(Spacer(1, 2 * mm))
            elif kind == "bullets":
                story.extend(pdf_bullets(value))
            elif kind == "steps":
                story.append(pdf_steps(value))
                story.append(Spacer(1, 2 * mm))
            elif kind == "cards":
                cells = []
                for label_text, value_text in value:
                    cell = Table(
                        [[Paragraph(label_text.upper(), pdf_styles_cache["label"])],
                         [Paragraph(value_text, pdf_styles_cache["cell"])]],
                        colWidths=[56 * mm],
                    )
                    cell.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, -1), PALE),
                        ("BOX", (0, 0), (-1, -1), 0.45, LINE),
                        ("LEFTPADDING", (0, 0), (-1, -1), 3 * mm),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 3 * mm),
                        ("TOPPADDING", (0, 0), (-1, -1), 2 * mm),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2 * mm),
                    ]))
                    cells.append(cell)
                grid = Table([cells], colWidths=[60 * mm] * 3)
                grid.setStyle(TableStyle([
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 0),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 1.5 * mm),
                ]))
                story.append(grid)
                story.append(Spacer(1, 2 * mm))
        if page_number < len(PAGES):
            story.append(PageBreak())

    doc = SimpleDocTemplate(
        str(PDF_OUT), pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm,
        topMargin=13 * mm, bottomMargin=21 * mm,
        title="Automação do anúncio de navios no OpenPort - Visão gerencial e operacional",
        author="CODEBA · Material informativo",
        subject="Documento de acompanhamento da iniciativa de automação do anúncio de navios",
    )
    doc.build(story, onFirstPage=pdf_page_footer, onLaterPages=pdf_page_footer)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def docx_add_paragraph(doc, text, style="Normal", *, color=INK_HEX, bold=False, size=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    return p


def docx_add_rich_paragraph(doc, text, style="Normal", *, color=INK_HEX, size=10):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    for part in text.split("**"):
        run = p.add_run(part)
        run.bold = len(p.runs) % 2 == 0
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor.from_string(color)
    return p


def docx_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(1.72)
    table.columns[1].width = Inches(5.28)
    for left, right in rows:
        cells = table.add_row().cells
        cells[0].width = Inches(1.72)
        cells[1].width = Inches(5.28)
        for idx, text in enumerate((left, right)):
            cell = cells[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(BLUE_HEX if idx == 0 else INK_HEX)
            run.bold = idx == 0
            if idx == 0:
                set_cell_shading(cell, PALE_HEX)
    return table


def docx_callout(doc, heading, text, warning=False):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, AMBER_PALE_HEX if warning else GREEN_PALE_HEX)
    set_cell_margins(cell, top=160, start=220, bottom=160, end=220)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(heading + "\n")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(AMBER_HEX if warning else GREEN_HEX)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(INK_HEX)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build_docx():
    document = Document()
    section_obj = document.sections[0]
    section_obj.top_margin = Inches(0.65)
    section_obj.bottom_margin = Inches(0.7)
    section_obj.left_margin = Inches(0.75)
    section_obj.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK_HEX)
    for heading_name, size, color in (
        ("Heading 1", 20, NAVY_HEX),
        ("Heading 2", 13, NAVY_HEX),
        ("Heading 3", 11, BLUE_HEX),
    ):
        style = document.styles[heading_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(5)

    header = section_obj.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("CODEBA  |  GPI / ÁREA OPERACIONAL")
    header_run.font.name = "Arial"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor.from_string(MUTED_HEX)
    footer = section_obj.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run(f"Documento informativo para discussão gerencial  ·  {AS_OF}")
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED_HEX)

    for page_number, page in enumerate(PAGES, start=1):
        if page_number == 1:
            brand = document.add_table(rows=1, cols=2)
            brand.autofit = False
            brand.columns[0].width = Inches(2.25)
            brand.columns[1].width = Inches(4.75)
            logo_cell, title_cell = brand.rows[0].cells
            logo_cell.width = Inches(2.25)
            title_cell.width = Inches(4.75)
            for cell in (logo_cell, title_cell):
                set_cell_shading(cell, NAVY_HEX)
                set_cell_margins(cell, top=170, start=200, bottom=170, end=200)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            logo_paragraph = logo_cell.paragraphs[0]
            logo_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if LOGO.exists():
                logo_paragraph.add_run().add_picture(str(LOGO), width=Inches(1.95))
            title_paragraph = title_cell.paragraphs[0]
            title_run = title_paragraph.add_run("GESTÃO PORTUÁRIA · GPI")
            title_run.bold = True
            title_run.font.name = "Arial"
            title_run.font.size = Pt(11)
            title_run.font.color.rgb = RGBColor.from_string("FFFFFF")
            document.add_paragraph().paragraph_format.space_after = Pt(1)
            p = document.add_paragraph()
            p.paragraph_format.space_before = Pt(22)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("AUTOMAÇÃO DO ANÚNCIO\nDE NAVIOS NO OPENPORT")
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(28)
            run.font.color.rgb = RGBColor.from_string(NAVY_HEX)
            docx_add_paragraph(document, "Visão gerencial e operacional · Documento informativo para discussão",
                               color=MUTED_HEX, size=12)
            docx_add_paragraph(document, AS_OF, color=MUTED_HEX, size=9)
        else:
            document.add_page_break()
            docx_add_paragraph(document, page["eyebrow"].upper(), color=BLUE_HEX, bold=True, size=8)
            docx_add_paragraph(document, page["title"], color=NAVY_HEX, bold=True, size=20)

        for kind, value in page["blocks"]:
            if kind == "heading":
                document.add_heading(value, level=2)
            elif kind in ("body", "lead", "body_small"):
                size = 11 if kind == "lead" else (8 if kind == "body_small" else 10)
                docx_add_paragraph(document, value, color=MUTED_HEX if kind == "body_small" else INK_HEX, size=size)
            elif kind in ("callout", "callout_warning"):
                docx_callout(document, value[0], value[1], warning=kind == "callout_warning")
            elif kind == "table":
                docx_table(document, value)
                document.add_paragraph().paragraph_format.space_after = Pt(1)
            elif kind == "bullets":
                for item in value:
                    p = document.add_paragraph(style="List Bullet")
                    p.paragraph_format.space_after = Pt(3)
                    run = p.add_run(item)
                    run.font.name = "Arial"
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor.from_string(INK_HEX)
            elif kind == "steps":
                docx_table(document, [
                    (f"{number} · {heading}", detail)
                    for number, heading, detail in value
                ])
            elif kind == "cards":
                docx_table(document, value)

    document.core_properties.title = "Automação do anúncio de navios no OpenPort - visão gerencial"
    document.core_properties.subject = "Iniciativa operacional da CODEBA"
    document.core_properties.author = "CODEBA · Material informativo"
    document.save(DOCX_OUT)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_pdf()
    build_docx()
    print(f"PDF: {PDF_OUT} ({PDF_OUT.stat().st_size:,} bytes)")
    print(f"DOCX: {DOCX_OUT} ({DOCX_OUT.stat().st_size:,} bytes)")


if __name__ == "__main__":
    main()
