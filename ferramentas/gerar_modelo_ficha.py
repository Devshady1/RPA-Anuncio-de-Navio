# -*- coding: utf-8 -*-
"""Gera o modelo da ficha unica a partir de ficha_anuncio_navio.xlsx.

O modelo preserva a aba Anuncio (inclusive a formula da C3), mescla C3:E3
para a mensagem de STATUS, atualiza as instrucoes de salvar/fechar e cria a
aba HISTORICO com o cabecalho das fichas processadas.

Uso:
    python ferramentas/gerar_modelo_ficha.py
"""
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
ORIGEM = ROOT / "ficha_anuncio_navio.xlsx"
DESTINO = ROOT / "src" / "anuncio_navio" / "recursos" / "ficha_modelo.xlsx"

CABECALHO_HISTORICO = (
    "data", "navio", "imo", "viagem", "duv", "resultado", "detalhe",
)
INSTRUCAO_FINAL = (
    "\u2714  Preencheu? Salve (Cmd+S) e FECHE \u2014 a ficha nova aparece sozinha."
)
INSTRUCAO_A9 = (
    "Ao terminar: salve (Cmd+S) e feche esta planilha. O rob\u00f4 processa o "
    "arquivo automaticamente; n\u00e3o \u00e9 necess\u00e1rio devolv\u00ea-lo."
)
INSTRUCAO_A10 = (
    "Em caso de erro, o motivo aparece na linha STATUS. Corrija os campos "
    "indicados e salve novamente."
)


def gerar(origem: Path, destino: Path) -> Path:
    """Aplica as edicoes do modelo e grava em `destino`."""
    wb = openpyxl.load_workbook(origem)
    ws = wb["An\u00fancio"]

    ws.merge_cells("C3:E3")
    ws["A29"] = INSTRUCAO_FINAL

    como = wb["COMO PREENCHER"]
    como["A9"] = INSTRUCAO_A9
    como["A10"] = INSTRUCAO_A10
    como["A10"].font = copy(como["A9"].font)
    como["A10"].alignment = copy(como["A9"].alignment)

    if "HIST\u00d3RICO" in wb.sheetnames:
        del wb["HIST\u00d3RICO"]
    hist = wb.create_sheet("HIST\u00d3RICO")
    hist.append(CABECALHO_HISTORICO)
    for cell in hist[1]:
        cell.font = Font(bold=True)

    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)
    return destino


def main() -> None:
    destino = gerar(ORIGEM, DESTINO)
    print(f"Modelo gerado: {destino}")


if __name__ == "__main__":
    main()
